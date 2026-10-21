"""Rapport Excel : recalculées par LibreOffice, les formules de l'onglet « Synthèse » redonnent les valeurs du moteur.

Ignoré si LibreOffice (soffice) n'est pas installé.
"""
import shutil
import subprocess
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
BASE = Path(__file__).resolve().parents[1] / "data" / "btp_fictif.duckdb"


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice absent")
def test_formules_excel_egales_au_moteur(tmp_path):
    if not BASE.exists():
        pytest.skip("base fictive absente")
    import openpyxl
    from qdd_s2.campagne import Config, lancer_campagnes
    from qdd_s2.rapport_excel import ecrire_rapport
    camps, var = lancer_campagnes(Config(base=BASE))
    source = tmp_path / "rapport.xlsx"
    ecrire_rapport(camps, var, source)
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir",
                    str(tmp_path / "recalcule"), str(source)], check=True, capture_output=True, timeout=180)
    ws = openpyxl.load_workbook(tmp_path / "recalcule" / "rapport.xlsx", data_only=True)["Synthèse"]
    comparees = 0
    for r in range(4, 20):
        excel, moteur = ws.cell(r, 2).value, ws.cell(r, 3).value
        if isinstance(excel, (int, float)) and isinstance(moteur, (int, float)):
            comparees += 1
            assert excel == pytest.approx(moteur, abs=1e-9), ws.cell(r, 1).value
    assert comparees == 11
